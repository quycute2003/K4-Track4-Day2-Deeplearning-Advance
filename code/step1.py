"""Backbone screening with the saved B01 recipe; no test inference."""
from __future__ import annotations

import gc
import json
import time
import zipfile
from dataclasses import asdict, fields
from pathlib import Path

import pandas as pd
import torch

if __package__:
    from . import model as models, train
else:
    import model as models
    import train

BACKBONES = [('B01', 'resnet50', 'ResNet'),
             ('B02', 'resnext50_32x4d', 'ResNeXt'),
             ('B03', 'convnext_tiny', 'ConvNeXt'),
             ('B04', 'deit_small_patch16_224', 'Transformer'),
             ('B05', 'efficientnet_b0', 'Lightweight')]
PATH_FIELDS = {'images_dir', 'labels_dir', 'out_dir', 'pred_dir', 'curves_dir'}
IGNORE_FIELDS = PATH_FIELDS | {'exp_id', 'backbone', 'device', 'resume'}
COLUMNS = ['exp_id', 'backbone', 'family', 'pretrained_tag', 'seed', 'epochs',
           'batch_size', 'params_m', 'gmacs', 'macro_f1_val', 'top1_val',
           'best_epoch', 'train_seconds_epoch', 'train_val_seconds_epoch',
           'latency_ms_batch1_fp32', 'gpu', 'status', 'config_path', 'curve_path']


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def baseline_config(project):
    record = read_json(Path(project) / 'runs/B01/seed0/config.json')
    cfg = record['config']
    required = dict(exp_id='B01', backbone='resnet50', seed=0, fold=0,
                    epochs=10, batch_size=32, img_size=224, init='finetune',
                    aug='basic', mix=None, loss='ce', save_test_predictions=False)
    if any(cfg.get(k) != v for k, v in required.items()):
        raise ValueError('B01 does not match the completed Step 0 baseline; do not mix recipes.')
    return record


def assert_same_recipe(reference, candidate):
    keys = {f.name for f in fields(train.Config)} - IGNORE_FIELDS
    changes = {k: (reference.get(k), candidate.get(k)) for k in keys
               if reference.get(k) != candidate.get(k)}
    if changes:
        raise ValueError(f'T00 recipe mismatch: {changes}')
    if candidate.get('save_test_predictions'):
        raise ValueError('Step 1 must not evaluate test.')


def assert_environment(reference):
    import timm, torchvision
    source = reference['config']['exp_id']
    if not torch.cuda.is_available():
        raise RuntimeError('Choose a T4 GPU runtime for screening.')
    actual = {'torch': torch.__version__, 'torchvision': torchvision.__version__,
              'timm': timm.__version__}
    for name, version in actual.items():
        if version != reference['versions'][name]:
            raise RuntimeError(f'{name}: current {version}, {source} used {reference["versions"][name]}. '
                               f'Use the same environment before comparing with {source}.')
    if torch.cuda.get_device_name(0) != reference['gpu']:
        raise RuntimeError(f'{source} used {reference["gpu"]}; choose the same GPU for timing comparisons.')


def make_config(reference, project, images_dir, labels_dir, exp_id, backbone):
    values = dict(reference['config'])
    root = Path(project)
    values.update(exp_id=exp_id, backbone=backbone, images_dir=str(images_dir),
                  labels_dir=str(labels_dir), out_dir=str(root / 'runs'),
                  pred_dir=str(root / 'predictions'), curves_dir=str(root / 'curves'),
                  device='auto', resume=False, save_test_predictions=False)
    assert_same_recipe(reference['config'], values)
    return train.Config(**values)


def import_baseline(project, archive, checkpoint):
    """Import only B01 evidence, preserving current implementation and CSVs."""
    root = Path(project).resolve()
    prefixes = ('runs/B01/seed0/', 'curves/B01_', 'predictions/B01_seed0_val.csv')
    with zipfile.ZipFile(archive) as bundle:
        if bundle.testzip() is not None:
            raise ValueError('Corrupt Step 0 archive')
        pending = []
        for info in bundle.infolist():
            target = (root / info.filename).resolve()
            if not target.is_relative_to(root):
                raise ValueError('Archive path escapes project')
            if info.is_dir() or not info.filename.startswith(prefixes):
                continue
            content = bundle.read(info)
            if target.exists() and target.read_bytes() != content:
                raise FileExistsError(f'Existing B01 evidence differs: {target}')
            pending.append((target, content))
        for target, content in pending:
            target.parent.mkdir(parents=True, exist_ok=True)
            if not target.exists():
                target.write_bytes(content)
    import shutil
    target = root / 'runs/B01/seed0/best.pt'
    if not target.exists():
        shutil.copyfile(checkpoint, target)
    baseline_config(root)
    return target


@torch.inference_mode()
def measure_latency(project, cfg, *, warmup=10):
    """One preliminary batch-1 FP32 forward after warmup; synchronized wall time.

    Excludes reading images, preprocessing, CPU/GPU transfers and model loading.
    This is not the p50/p95/p99 benchmark required in Step 3.
    """
    folder = train.run_dir(cfg)
    record = read_json(folder / 'config.json')
    pretrained = record['pretrained_cfg']
    name = pretrained['architecture'] + '.' + pretrained['tag']
    train.set_seed(cfg.seed)
    net = models.build_model(name, pretrained=False, init='scratch',
                             img_size=cfg.img_size).eval().cuda()
    state = torch.load(folder / 'best.pt', map_location='cpu', weights_only=True)
    summary = read_json(folder / 'summary.json')
    if (state.get('epoch') != summary['best_epoch'] or state.get('macro_f1') is None
            or abs(state['macro_f1'] - summary['macro_f1_val']) > 1e-8):
        raise ValueError('best.pt is not the checkpoint named in the summary')
    net.load_state_dict(state['model'], strict=True)
    x = torch.zeros(1, 3, cfg.img_size, cfg.img_size, device='cuda')
    output = None
    try:
        for _ in range(warmup):
            net(x)
        torch.cuda.synchronize()
        start = time.perf_counter()
        output = net(x)
        torch.cuda.synchronize()
        milliseconds = (time.perf_counter() - start) * 1000
        if output.shape != (1, 9) or not torch.isfinite(output).all():
            raise ValueError('Invalid latency forward output')
        result = dict(latency_ms=milliseconds, batch_size=1, img_size=cfg.img_size,
                      dtype='FP32', warmup=warmup, measured_forwards=1,
                      scope='model forward only; synchronized wall time; synthetic zero input',
                      gpu=torch.cuda.get_device_name(0), torch=torch.__version__,
                      preliminary=True, checkpoint_epoch=state['epoch'])
        train.write_json(folder / 'latency_preliminary.json', result)
        return result
    finally:
        del net, state, x, output
        gc.collect()
        torch.cuda.empty_cache()


def result_row(project, exp_id, backbone, family, reference):
    root = Path(project)
    folder = root / 'runs' / exp_id / 'seed0'
    row = dict.fromkeys(COLUMNS)
    row.update(exp_id=exp_id, backbone=backbone, family=family, seed=0,
               epochs=reference['config']['epochs'], batch_size=reference['config']['batch_size'],
               status='Chưa chạy', config_path=f'runs/{exp_id}/seed0/config.json',
               curve_path=f'curves/{exp_id}_{backbone}_seed0.png')
    if not (folder / 'summary.json').exists():
        return row
    record, summary = read_json(folder / 'config.json'), read_json(folder / 'summary.json')
    assert_same_recipe(reference['config'], record['config'])
    if (summary['test_evaluated'] or summary['backbone'] != backbone
            or record['config']['backbone'] != backbone or summary['exp_id'] != exp_id):
        raise ValueError('Invalid screening result')
    if record['gpu'] != reference['gpu'] or any(
            record['versions'][key] != reference['versions'][key]
            for key in ('torch', 'torchvision', 'timm')):
        raise ValueError('Training environment differs from B01')
    history = pd.read_csv(folder / 'history.csv')
    if len(history) != row['epochs']:
        raise ValueError('Incomplete training history')
    if not (root / row['curve_path']).is_file():
        raise FileNotFoundError(f'Missing training curve: {row["curve_path"]}')
    row.update(pretrained_tag=record['pretrained_cfg']['tag'], params_m=summary['params_m'],
               gmacs=summary['gmacs'], macro_f1_val=summary['macro_f1_val'],
               top1_val=summary['top1_val'], best_epoch=summary['best_epoch'],
               train_seconds_epoch=float(history.train_seconds.mean()),
               train_val_seconds_epoch=float(history.epoch_seconds.mean()), gpu=record['gpu'],
               status='Đã train; chưa đo latency')
    latency_path = folder / 'latency_preliminary.json'
    if latency_path.exists():
        latency = read_json(latency_path)
        if latency['gpu'] != reference['gpu'] or latency['torch'] != reference['versions']['torch']:
            raise ValueError('Latency was measured in a different environment')
        row.update(latency_ms_batch1_fp32=latency['latency_ms'], status='Đủ kết quả')
    return row


def collect_results(project):
    root = Path(project)
    reference = baseline_config(root)
    rows = [result_row(root, *item, reference) for item in BACKBONES]
    out = root / 'step1'
    out.mkdir(parents=True, exist_ok=True)
    train.write_json(out / 'backbones.json', rows)
    pd.DataFrame(rows, columns=COLUMNS).to_csv(out / 'backbones.csv', index=False)
    return rows


def capture_imagenet_reference(project):
    """Match exact pretrained tags at 224; never substitute another tag."""
    import hashlib
    import io
    import urllib.request
    from datetime import datetime, timezone
    url = 'https://raw.githubusercontent.com/huggingface/pytorch-image-models/main/results/results-imagenet.csv'
    content = urllib.request.urlopen(url, timeout=30).read()
    reference = pd.read_csv(io.BytesIO(content))
    rows = collect_results(project)
    names = [row['backbone'] + '.' + row['pretrained_tag'] for row in rows
             if row['pretrained_tag']]
    matched = reference[(reference.model.isin(names)) & (reference.img_size == 224)]
    if matched.model.duplicated().any():
        raise ValueError('Ambiguous ImageNet reference rows')
    out = Path(project) / 'step1'
    matched.to_csv(out / 'imagenet_reference.csv', index=False)
    train.write_json(out / 'imagenet_source.json', dict(url=url,
                     fetched_utc=datetime.now(timezone.utc).isoformat(),
                     source_sha256=hashlib.sha256(content).hexdigest(),
                     matching='exact architecture.tag, img_size=224',
                     missing=[name for name in names if name not in set(matched.model)]))
    return matched


def run_backbone(project, images_dir, labels_dir, exp_id):
    """Resume only the same recipe; completed runs are validated and reused."""
    reference = baseline_config(project)
    assert_environment(reference)
    _, backbone, _ = next(item for item in BACKBONES if item[0] == exp_id)
    cfg = make_config(reference, project, images_dir, labels_dir, exp_id, backbone)
    folder = train.run_dir(cfg)
    if (folder / 'summary.json').exists():
        existing = read_json(folder / 'config.json')
        assert_same_recipe(reference['config'], existing['config'])
        if existing['config']['backbone'] != backbone:
            raise ValueError('Existing experiment uses a different backbone')
        result_row(project, exp_id, backbone, next(i[2] for i in BACKBONES if i[0] == exp_id), reference)
        print(f'{exp_id}: reuse completed result; no retraining.', flush=True)
    elif exp_id == 'B01':
        raise FileNotFoundError('Import the completed B01 results and best.pt first.')
    else:
        cfg.resume = (folder / 'last.pt').exists()
        if cfg.resume:
            existing = read_json(folder / 'config.json')
            assert_same_recipe(reference['config'], existing['config'])
            if existing['config']['backbone'] != backbone:
                raise ValueError('Cannot resume a different backbone')
            # The saved Colab paths must match for exact train.run resume checks.
        train.run(cfg)
    measure_latency(project, cfg)
    collect_results(project)
    return read_json(folder / 'summary.json')


def write_report(project):
    root = Path(project)
    rows = collect_results(root)
    complete = [r for r in rows if r['status'] == 'Đủ kết quả']
    if len(complete) != 5:
        raise ValueError('Finish all five backbones and latency measurements before selection.')
    ranked = sorted(complete, key=lambda r: (-r['macro_f1_val'], r['latency_ms_batch1_fp32']))
    best, fastest = ranked[0], min(complete, key=lambda r: r['latency_ms_batch1_fp32'])
    selection = dict(selected_exp_ids=[best['exp_id']], backbone=best['backbone'],
                     recipe='T00', seed=0, rule='highest val macro-F1; latency breaks exact ties',
                     rationale=f"{best['exp_id']} has highest val macro-F1 ({best['macro_f1_val']:.6f}); "
                               'one backbone for Step 2 under the recorded GPU budget.',
                     test_used=False, single_seed=True)
    train.write_json(root / 'step1/selection.json', selection)
    lines = ['# Bước 1 — So sánh backbone', '',
             'T00: kế thừa cấu hình B01, seed 0, fold 0, 10 epoch, batch 32, input 224, '
             'finetune ImageNet, AdamW, LR backbone/head 1e-4/1e-3, WD 0,05 '
             '(trừ norm/bias), warmup 1 epoch, cosine, CE, AMP. Chọn checkpoint bằng macro-F1 val.', '',
             'Mean/std và nội suy lấy từ pretrained_cfg của từng tag; resize val 256 + CenterCrop 224. '
             'Đây là so sánh kiến trúc cùng công thức finetune, còn các tag ImageNet có lịch sử pretrain khác nhau.', '',
             '| exp_id | Backbone | Tag | Params (M) | GMAC | F1 val | Top-1 val | Train s/epoch | Latency ms |',
             '|---|---|---|---:|---:|---:|---:|---:|---:|']
    for r in rows:
        lines.append(f"| {r['exp_id']} | {r['backbone']} | {r['pretrained_tag']} | "
                     f"{r['params_m']:.3f} | {r['gmacs']:.3f} | {r['macro_f1_val']:.2%} | "
                     f"{r['top1_val']:.2%} | {r['train_seconds_epoch']:.2f} | {r['latency_ms_batch1_fp32']:.3f} |")
    lines += ['', f"Chọn **{best['exp_id']} — {best['backbone']}** cho Bước 2, "
              f"macro-F1 val **{best['macro_f1_val']:.2%}**, latency sơ bộ "
              f"**{best['latency_ms_batch1_fp32']:.3f} ms**. Kế hoạch GPU dành ablation cho một backbone.",
              f"Nhanh nhất trong phép đo sơ bộ: {fastest['exp_id']} ({fastest['latency_ms_batch1_fp32']:.3f} ms), "
              f"macro-F1 val {fastest['macro_f1_val']:.2%}.", '',
              f"Chênh F1 giữa hai model đứng đầu: {(ranked[0]['macro_f1_val']-ranked[1]['macro_f1_val'])*100:.3f} điểm phần trăm. "
              'Mỗi model chỉ có một seed: thứ hạng này là quan sát screening, chưa chứng minh khác biệt thống kê.', '',
              '## Hội tụ và dấu hiệu overfit', '']
    for r in rows:
        h = pd.read_csv(root / 'runs' / r['exp_id'] / 'seed0/history.csv')
        threshold = .95 * h.val_macro_f1.max()
        first = int(h.loc[h.val_macro_f1 >= threshold, 'epoch'].iloc[0])
        end = h.tail(3)
        signal = (end.train_loss.diff().dropna() < 0).all() and (end.val_loss.diff().dropna() > 0).all()
        lines.append(f"- {r['exp_id']}: đạt 95% F1 tốt nhất của chính model từ epoch {first}; "
                     f"train loss {h.iloc[0].train_loss:.4f} → {h.iloc[-1].train_loss:.4f}, "
                     f"val loss {h.iloc[0].val_loss:.4f} → {h.iloc[-1].val_loss:.4f}. "
                     + ('Ba epoch cuối có train loss giảm và val loss tăng liên tiếp; cần xem xét overfit.' if signal else
                        'Ba epoch cuối không có tín hiệu train loss giảm/val loss tăng liên tiếp; vẫn cần đọc toàn bộ đường cong.'))
    lines += ['', 'Mốc 95% là tiêu chí mô tả so với đỉnh riêng của mỗi model; '
              'không đồng nghĩa model đạt chất lượng tuyệt đối tốt hơn hoặc hội tụ hoàn toàn.', '',
              '## Tốc độ, GMAC và giới hạn', '',
              'Latency: batch 1, FP32, 10 warmup rồi đo một forward, synchronize CUDA hai đầu; '
              'input tensor 0 đã ở GPU, chỉ tính forward, loại tải ảnh/tiền xử lý/truyền dữ liệu. '
              'Cùng T4 và phiên bản thư viện. Đây là phép đo sơ bộ có nhiễu; p50/p95/p99 ở Bước 3.',
              'GMAC từ fvcore, 1 FMA = 1 operation; xem unsupported_ops trong config từng run. '
              'GMAC không đo truy cập bộ nhớ hay hiệu quả kernel; đối chiếu bảng trước khi dùng để suy ra tốc độ.', '']
    by_compute = sorted(rows, key=lambda r: r['gmacs'])
    by_train = sorted(rows, key=lambda r: r['train_seconds_epoch'])
    by_latency = sorted(rows, key=lambda r: r['latency_ms_batch1_fp32'])
    for label, order in [('GMAC tăng dần', by_compute), ('Thời gian train tăng dần', by_train),
                         ('Latency sơ bộ tăng dần', by_latency)]:
        lines.append(label + ': ' + ' → '.join(r['exp_id'] for r in order) + '.')
    same_order = [r['exp_id'] for r in by_compute] == [r['exp_id'] for r in by_latency]
    lines.append('Thứ tự GMAC và latency ' + ('trùng' if same_order else 'không trùng') +
                 ' trong phép đo này; một forward/model chưa đủ để đánh giá độ ổn định của tương quan.')
    lines += ['', '## Đối chiếu ImageNet', '']
    imagenet = root / 'step1/imagenet_reference.csv'
    source = root / 'step1/imagenet_source.json'
    if imagenet.exists() and source.exists():
        ref = pd.read_csv(imagenet)
        names = [r['backbone'] + '.' + r['pretrained_tag'] for r in rows]
        matched = ref[ref.model.isin(names) & (ref.img_size == 224)]
        lines += ['Nguồn: [bảng kết quả timm](' + read_json(source)['url'] + '), '
                  'khớp chính xác tag và input 224; snapshot trong imagenet_reference.csv.', '',
                  '| Tag | Top-1 ImageNet | Crop pct | Nội suy |', '|---|---:|---:|---|']
        for item in matched.itertuples():
            lines.append(f'| {item.model} | {item.top1:.3f}% | {item.crop_pct:.3f} | {item.interpolation} |')
        if len(matched) == 5:
            im_order = matched.sort_values('top1', ascending=False).model.tolist()
            dw_order = [r['backbone'] + '.' + r['pretrained_tag'] for r in ranked]
            lines += ['', 'ImageNet top-1 giảm dần: ' + ' → '.join(im_order) + '.',
                      'DeepWeeds macro-F1 giảm dần: ' + ' → '.join(dw_order) + '.',
                      'Hai thứ hạng ' + ('trùng nhau.' if im_order == dw_order else 'khác nhau.'),
                      'Hai bài toán/metric và điều kiện đánh giá khác nhau; top-1 ImageNet không bảo đảm '
                      'thứ hạng macro-F1 sau finetune trên DeepWeeds.']
        else:
            lines.append('Thiếu một số tag tương ứng; chưa kết luận thứ hạng đầy đủ.')
    else:
        lines.append('Chưa có snapshot ImageNet cho đúng các tag; không thay bằng accuracy của tag khác.')
    lines += ['', '## Bằng chứng', '']
    for r in rows:
        lines.append(f"![{r['exp_id']} curve](../{r['curve_path']})")
    lines += ['', 'Cấu hình/log/checkpoint: runs/B0x/seed0. Dự đoán val: predictions/B0x_seed0_val.csv. '
              'Không đánh giá test trong Bước 1.']
    (root / 'step1/step1_report.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    return selection


def export_colab_xlsx(project):
    return write_colab_sheet(project, collect_results(project), COLUMNS, 'Backbones', [
        'T00: 10 epoch, batch 32, seed 0. Một seed mỗi backbone; chỉ val.',
        'Latency sơ bộ: FP32, batch 1, 10 warmup + một forward; chỉ thời gian model trên GPU.',
        'Nguồn: runs/B0x/seed0/{config,history,summary,latency_preliminary}.json/csv.'])


def write_colab_sheet(project, rows, columns, sheet_name, notes):
    """Colab fallback: the desktop artifact-tool runtime is unavailable there.

    Preserve any other sheets; the desktop uses tools/export_backbones.mjs.
    """
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    path = Path(project) / 'results.xlsx'
    wb = load_workbook(path) if path.exists() else Workbook()
    index = wb.sheetnames.index(sheet_name) if sheet_name in wb else len(wb.sheetnames)
    if sheet_name in wb:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name, index)
    if 'Sheet' in wb and wb['Sheet'].max_row == 1 and wb['Sheet']['A1'].value is None:
        del wb['Sheet']
    ws.append(columns)
    for row in rows:
        ws.append([row[c] for c in columns])
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='243B53')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 42
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            cell.font = Font(name='Arial', size=10)
            cell.alignment = Alignment(vertical='center')
    from openpyxl.utils import get_column_letter
    for i, name in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = 22 if i != 2 else 30
        if name.endswith('_path'):
            ws.column_dimensions[get_column_letter(i)].width = 48
        for index in range(2, len(rows)+2):
            cell = ws.cell(index, i)
            if name in ('macro_f1_val', 'top1_val', 'balanced_acc_val') or name.startswith('f1_'):
                cell.number_format = '0.00%'
            elif name in ('params_m', 'gmacs', 'latency_ms_batch1_fp32'):
                cell.number_format = '0.000'
            elif name.endswith('seconds_epoch'):
                cell.number_format = '0.00'
            elif name == 'delta_f1_pp':
                cell.number_format = '+0.000;-0.000;0.000'
                metric_col = get_column_letter(columns.index('macro_f1_val')+1)
                cell.value = f'=IF(ISNUMBER({metric_col}{index}),100*({metric_col}{index}-${metric_col}$2),"")'
            elif name in ('ece_val', 'nll_val'):
                cell.number_format = '0.0000'
            elif name == 'relative_cost':
                metric_col = get_column_letter(columns.index('p50_ms')+1)
                cell.number_format = '0.00"×"'
                cell.value = f'=IF(AND(ISNUMBER({metric_col}{index}),ISNUMBER(${metric_col}$2),${metric_col}$2>0),{metric_col}{index}/${metric_col}$2,"")'
            elif name == 'p95_le_100ms':
                latency_col = get_column_letter(columns.index('p95_ms')+1)
                cell.value = f'=IF(ISNUMBER({latency_col}{index}),IF({latency_col}{index}<=100,"Đạt","Không đạt"),"Chờ đo")'
            elif name.endswith('_ms'):
                cell.number_format = '0.000'
            elif name.startswith('images_per_s'):
                cell.number_format = '0.0'
            elif name == 'temperature':
                cell.number_format = '0.000000'
    for index, note in enumerate(notes, len(rows)+4):
        ws.cell(index, 1, note)
    wb.save(path)
    return path


def export_artifacts(project, out_path):
    root = Path(project)
    with zipfile.ZipFile(out_path, 'w', zipfile.ZIP_DEFLATED) as bundle:
        for folder in ('code', 'tests', 'step0', 'step1', 'runs', 'curves', 'predictions'):
            for path in sorted((root / folder).rglob('*')):
                if path.is_file() and path.suffix in ('.py', '.md', '.json', '.csv', '.png', '.npy', '.ipynb'):
                    bundle.write(path, path.relative_to(root))
        for name in ('eval.py', 'results.xlsx'):
            bundle.write(root / name, name)
    return Path(out_path)
